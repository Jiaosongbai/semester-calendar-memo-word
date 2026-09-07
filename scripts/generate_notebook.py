#!/usr/bin/env python3
"""Generate an editable minimalist Word weekly planner from a validated JSON spec."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

CN = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十", "十一", "十二", "十三", "十四", "十五", "十六", "十七", "十八", "十九", "二十", "二十一", "二十二", "二十三", "二十四", "二十五"]
WEEKDAYS = ["周一", "周二", "周三", "周四", "周五", "周六", "周天"]
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def parse_date(value: str, field: str) -> date:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field} 必须是 YYYY-MM-DD 日期：{value!r}") from exc


def validate(spec: dict) -> None:
    for key in ("title", "sections"):
        if not spec.get(key):
            raise ValueError(f"缺少必填字段：{key}")
    if not isinstance(spec["sections"], list):
        raise ValueError("sections 必须是数组")
    for idx, section in enumerate(spec["sections"], 1):
        for key in ("name", "start", "weeks"):
            if section.get(key) in (None, ""):
                raise ValueError(f"sections[{idx}] 缺少字段：{key}")
        start = parse_date(section["start"], f"sections[{idx}].start")
        if start.weekday() != 0:
            raise ValueError(f"sections[{idx}].start 必须是星期一：{start}")
        if not isinstance(section["weeks"], int) or not 1 <= section["weeks"] <= 25:
            raise ValueError(f"sections[{idx}].weeks 必须是 1–25 的整数")
        exam = section.get("exam_from_week")
        if exam is not None and (not isinstance(exam, int) or not 1 <= exam <= section["weeks"]):
            raise ValueError(f"sections[{idx}].exam_from_week 超出范围")
    for group in ("statutory_holidays", "school_breaks"):
        for idx, item in enumerate(spec.get(group, []), 1):
            start = parse_date(item.get("start"), f"{group}[{idx}].start")
            end = parse_date(item.get("end"), f"{group}[{idx}].end")
            if end < start or not item.get("name"):
                raise ValueError(f"{group}[{idx}] 的名称或区间无效")
    for idx, item in enumerate(spec.get("calendar_day_markers", []), 1):
        parse_date(item.get("date"), f"calendar_day_markers[{idx}].date")
        if not item.get("name"):
            raise ValueError(f"calendar_day_markers[{idx}] 缺少 name")
    for idx, item in enumerate(spec.get("activities", []), 1):
        anchor = parse_date(item.get("anchor_date"), f"activities[{idx}].anchor_date")
        if item.get("precision") not in ("day", "month", "range") or not item.get("text"):
            raise ValueError(f"activities[{idx}] 的 precision 或 text 无效")
        covered = False
        for section in spec["sections"]:
            start = parse_date(section["start"], "section.start")
            if start <= anchor <= start + timedelta(days=section["weeks"] * 7 - 1):
                covered = True
                break
        if not covered:
            raise ValueError(f"activities[{idx}].anchor_date 不在任何生成区间内：{anchor}")


def set_run(run, *, bold=False, italic=False, size=12, font="宋体", color="222222", superscript=False):
    run.bold, run.italic = bold, italic
    run.font.size, run.font.name = Pt(size), font
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.superscript = superscript
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)


def keep_next(paragraph, enabled=True):
    ppr = paragraph._p.get_or_add_pPr()
    for node in ppr.findall(qn("w:keepNext")):
        ppr.remove(node)
    if enabled:
        ppr.append(OxmlElement("w:keepNext"))


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def cell_margins(cell, value=80):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.first_child_found_in("w:tcMar")
    if old is not None:
        tcpr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def add_cell_text(cell, text="", *, bold=False, italic=False, size=10, color="222222", marker_ids=()):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(text), bold=bold, italic=italic, size=size, color=color)
    for marker_id in marker_ids:
        set_run(p.add_run(f"[[FN{marker_id}]]"), size=14, color="777777", superscript=True)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    return p


def expand_ranges(items):
    result = {}
    for item in items:
        current = parse_date(item["start"], "range.start")
        end = parse_date(item["end"], "range.end")
        while current <= end:
            result.setdefault(current, []).append(item["name"])
            current += timedelta(days=1)
    return result


def safe_filename(title):
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip(" .")
    return (cleaned or "calendar-notebook") + ".docx"


def add_week_table(doc, monday, week_no, exam, day_notes, activities):
    table = doc.add_table(rows=4, cols=4)
    table.style, table.autofit = "Table Grid", False
    heights = (0.65, 1.95, 0.65, 1.95)
    for row, height in zip(table.rows, heights):
        no_split(row)
        row.height_rule, row.height = WD_ROW_HEIGHT_RULE.AT_LEAST, Cm(height)
        for cell in row.cells:
            cell.width = Cm(4.55)
    label = f"第{CN[week_no] if week_no < len(CN) else week_no}周" + ("（考试周）" if exam else "")
    add_cell_text(table.cell(0, 0), label, bold=True, color="FFFFFF")
    shade(table.cell(0, 0), "333333")
    add_cell_text(table.cell(1, 0))
    days = [monday + timedelta(days=i) for i in range(7)]
    for offset, current in enumerate(days):
        if offset < 3:
            row_label, row_body, col = 0, 1, offset + 1
        else:
            row_label, row_body, col = 2, 3, offset - 3
        stamp = f"{current.year}/{current.month}/{current.day}" if offset == 0 else f"{current.month}/{current.day}"
        notes = day_notes.get(current, [])
        ids = [item["id"] for item in activities.get(current, [])]
        add_cell_text(table.cell(row_label, col), f"{WEEKDAYS[offset]}·{stamp}", bold=True, italic=bool(notes), marker_ids=ids)
        add_cell_text(table.cell(row_body, col), "\n".join(notes), italic=bool(notes), size=9, color="8A3B2E" if notes else "222222")
    for p in [p for row in table.rows for cell in row.cells for p in cell.paragraphs]:
        keep_next(p, True)
    for cell in table.rows[3].cells:
        for p in cell.paragraphs:
            keep_next(p, False)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(5)
    spacer.paragraph_format.line_spacing = Pt(5)


def inject_footnotes(docx_path, activities):
    temp = Path(tempfile.mkdtemp(prefix="calendar-docx-"))
    try:
        with zipfile.ZipFile(docx_path) as archive:
            archive.extractall(temp)
        document_xml = temp / "word/document.xml"
        tree = etree.parse(str(document_xml))
        ns = {"w": W_NS}
        for item in activities:
            nodes = tree.xpath(f'//w:t[text()="[[FN{item["id"]}]]"]', namespaces=ns)
            if len(nodes) != 1:
                raise RuntimeError(f"脚注 {item['id']} 的正文锚点数量不是 1")
            run = nodes[0].getparent()
            for child in list(run):
                run.remove(child)
            rpr = etree.SubElement(run, f"{{{W_NS}}}rPr")
            for tag, value in (("color", "777777"), ("sz", "28"), ("szCs", "28"), ("vertAlign", "superscript")):
                node = etree.SubElement(rpr, f"{{{W_NS}}}{tag}")
                node.set(f"{{{W_NS}}}val", value)
            ref = etree.SubElement(run, f"{{{W_NS}}}footnoteReference")
            ref.set(f"{{{W_NS}}}id", str(item["id"]))
        tree.write(str(document_xml), xml_declaration=True, encoding="UTF-8", standalone=True)

        parts = [
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>',
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>',
        ]
        for item in activities:
            parts.append(
                f'<w:footnote w:id="{item["id"]}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
                '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/><w:color w:val="777777"/><w:sz w:val="28"/><w:szCs w:val="28"/><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
                '<w:r><w:rPr><w:color w:val="777777"/><w:sz w:val="16"/><w:szCs w:val="16"/><w:rFonts w:ascii="SimSun" w:hAnsi="SimSun" w:eastAsia="宋体"/></w:rPr>'
                f'<w:t xml:space="preserve">  {escape(item["text"])}</w:t></w:r></w:p></w:footnote>'
            )
        (temp / "word/footnotes.xml").write_text('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="' + W_NS + '">' + "".join(parts) + "</w:footnotes>", encoding="utf-8")

        styles_path = temp / "word/styles.xml"
        styles = etree.parse(str(styles_path))
        found = styles.xpath('//w:style[@w:styleId="FootnoteReference"]', namespaces=ns)
        style = found[0] if found else etree.SubElement(styles.getroot(), f"{{{W_NS}}}style", {f"{{{W_NS}}}type": "character", f"{{{W_NS}}}styleId": "FootnoteReference"})
        rpr = style.find(f"{{{W_NS}}}rPr") or etree.SubElement(style, f"{{{W_NS}}}rPr")
        for tag in ("color", "sz", "szCs", "vertAlign"):
            for node in rpr.findall(f"{{{W_NS}}}{tag}"):
                rpr.remove(node)
        for tag, value in (("color", "777777"), ("sz", "28"), ("szCs", "28"), ("vertAlign", "superscript")):
            node = etree.SubElement(rpr, f"{{{W_NS}}}{tag}")
            node.set(f"{{{W_NS}}}val", value)
        styles.write(str(styles_path), xml_declaration=True, encoding="UTF-8", standalone=True)

        rels = temp / "word/_rels/document.xml.rels"
        rel_tree = etree.parse(str(rels))
        rel_ns = rel_tree.getroot().nsmap.get(None)
        if not rel_tree.xpath('//*[local-name()="Relationship" and contains(@Type,"/footnotes")]'):
            etree.SubElement(rel_tree.getroot(), f"{{{rel_ns}}}Relationship", Id="rIdFootnotes", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes", Target="footnotes.xml")
        rel_tree.write(str(rels), xml_declaration=True, encoding="UTF-8", standalone=True)
        content_types = temp / "[Content_Types].xml"
        ct_tree = etree.parse(str(content_types))
        ct_ns = ct_tree.getroot().nsmap.get(None)
        if not ct_tree.xpath('//*[local-name()="Override" and @PartName="/word/footnotes.xml"]'):
            etree.SubElement(ct_tree.getroot(), f"{{{ct_ns}}}Override", PartName="/word/footnotes.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
        ct_tree.write(str(content_types), xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in temp.rglob("*"):
                if path.is_file():
                    archive.write(path, path.relative_to(temp))
    finally:
        shutil.rmtree(temp)


def generate(spec, output):
    validate(spec)
    official = expand_ranges(spec.get("statutory_holidays", []))
    breaks = expand_ranges(spec.get("school_breaks", []))
    markers = {parse_date(x["date"], "marker.date"): x["name"] for x in spec.get("calendar_day_markers", [])}
    day_notes = {}
    for current, names in official.items():
        day_notes.setdefault(current, []).extend(f"{name}（法定放假）" for name in names)
    for current, name in markers.items():
        day_notes.setdefault(current, []).append(f"{name}（校历标记）")
    for current, names in breaks.items():
        day_notes.setdefault(current, []).extend(f"{name}（学校假期）" for name in names)
    ordered = sorted(spec.get("activities", []), key=lambda x: (x["anchor_date"], x.get("duration_rank", 1), x["text"]))
    for idx, item in enumerate(ordered, 1):
        item["id"] = idx
    activity_map = {}
    for item in ordered:
        activity_map.setdefault(parse_date(item["anchor_date"], "activity.anchor_date"), []).append(item)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.25), Cm(1.2)
    sec.left_margin = sec.right_margin = Cm(1.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "宋体", Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    title = doc.add_paragraph()
    title.alignment, title.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(2)
    set_run(title.add_run(spec["title"]), bold=True, size=14, font="黑体", color="111111")
    subtitle = doc.add_paragraph()
    subtitle.alignment, subtitle.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(8)
    set_run(subtitle.add_run(spec.get("subtitle", "")), size=10, color="777777")
    for section_no, section in enumerate(spec["sections"]):
        if section_no:
            doc.add_page_break()
        heading = doc.add_paragraph()
        heading.alignment, heading.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(3)
        keep_next(heading)
        set_run(heading.add_run(section["name"]), bold=True, size=12, font="黑体")
        if section.get("info"):
            info = doc.add_paragraph()
            info.alignment, info.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(5)
            keep_next(info)
            set_run(info.add_run(section["info"]), size=8, color="666666")
        start = parse_date(section["start"], "section.start")
        for week in range(1, section["weeks"] + 1):
            monday = start + timedelta(days=7 * (week - 1))
            exam = bool(section.get("exam_from_week") and week >= section["exam_from_week"])
            add_week_table(doc, monday, week, exam, day_notes, activity_map)
        if spec.get("source_note"):
            note = doc.add_paragraph(spec["source_note"])
            note.paragraph_format.space_before, note.paragraph_format.space_after = Pt(4), Pt(0)
            for run in note.runs:
                set_run(run, size=7, color="777777")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    if ordered:
        inject_footnotes(output, ordered)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec", type=Path, help="符合 references/spec-schema.md 的 JSON 文件")
    parser.add_argument("-o", "--output", type=Path, help="输出 DOCX 路径；默认写入当前目录")
    args = parser.parse_args()
    try:
        spec = json.loads(args.spec.read_text(encoding="utf-8"))
        output = args.output or Path.cwd() / safe_filename(spec.get("title", "calendar-notebook"))
        print(generate(spec, output).resolve())
    except (OSError, json.JSONDecodeError, ValueError, RuntimeError) as exc:
        print(f"无法生成：{exc}", file=sys.stderr)
        raise SystemExit(2) from exc


if __name__ == "__main__":
    main()
